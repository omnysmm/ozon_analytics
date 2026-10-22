import openpyxl
import os

files = [
    r"C:\Users\Alex\Downloads\Telegram Desktop\pugovits yassı (2).xlsx",
    r"C:\Users\Alex\Yandex.Disk\Pugovicnet\Цены фурнитура ВБ.xlsx",
    r"C:\Users\Alex\Yandex.Disk\Pugovicnet\прайс для 1С.xlsx",
    r"C:\Users\Alex\Yandex.Disk\Загрузки\Копия запрос по ценам3 .xlsx",
    r"C:\Users\Alex\Yandex.Disk\Pugovicnet\FBS\stocks press strong.xlsx",
]

for fp in files:
    print("=" * 80)
    print(fp, "exists=", os.path.exists(fp))
    if not os.path.exists(fp):
        continue
    try:
        wb = openpyxl.load_workbook(fp, data_only=True, read_only=True)
    except Exception as e:
        print("ERR", e)
        continue
    print("sheets:", wb.sheetnames)
    for sn in wb.sheetnames[:4]:
        ws = wb[sn]
        print(f"--- sheet {sn} dims={ws.dimensions if hasattr(ws,'dimensions') else '?'} ---")
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            rows.append(row)
            if i >= 12:
                break
        for i, r in enumerate(rows):
            vals = [str(c)[:40] if c is not None else "" for c in r[:18]]
            print(i, " | ".join(vals))
    wb.close()
