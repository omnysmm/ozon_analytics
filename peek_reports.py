import openpyxl
import os

files = [
    r"C:\Users\Alex\Downloads\Отчет о реализации товара_20260930.xlsx",
    r"C:\Users\Alex\Downloads\Отчет по начислениям_14.09.2026-14.10.2026.xlsx",
    r"C:\Users\Alex\Downloads\Отчет по логистике за период 2026-09-14 - 2026-10-14.xlsx",
    r"C:\Users\Alex\Downloads\Отчёт по кэшбэкам 2026-08-03_2026-08-09.xlsx",
    r"C:\Users\Alex\Downloads\Еженедельный_отчет_2024_01_29_2026_04_19_112393.xlsx",
    r"C:\Users\Alex\Downloads\Отчет о проблемных товарах 06.04-12.04.xlsx",
]

for fp in files:
    print("=" * 90)
    print(os.path.basename(fp), os.path.getsize(fp))
    if not os.path.exists(fp):
        print("MISSING")
        continue
    try:
        wb = openpyxl.load_workbook(fp, data_only=True, read_only=True)
    except Exception as e:
        print("ERR", type(e).__name__, e)
        continue
    print("sheets:", wb.sheetnames)
    for sn in wb.sheetnames[:3]:
        ws = wb[sn]
        print(f"--- {sn} ---")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i > 8:
                break
            vals = [str(c)[:35] if c is not None else "" for c in row[:20]]
            print(i, " | ".join(vals))
    wb.close()
